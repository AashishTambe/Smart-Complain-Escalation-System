import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    # Add code as monospaced paragraphs (approximation).
    for line in code_lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(10)


def convert(md_path: Path, docx_path: Path) -> None:
    doc = Document()

    md_text = md_path.read_text(encoding="utf-8")
    lines = md_text.splitlines()

    in_code_block = False
    code_fence_lang = None
    code_lines: list[str] = []

    def flush_code_block():
        nonlocal in_code_block, code_lines, code_fence_lang
        if code_lines:
            add_code_block(doc, code_lines)
        in_code_block = False
        code_fence_lang = None
        code_lines = []

    # Title from first H1 if present
    for line in lines:
        if line.startswith("# "):
            title = line[2:].strip()
            if title:
                doc.add_heading(title, level=0)
            break

    # Start from after that first H1 line
    start_idx = 0
    for i, line in enumerate(lines):
        if line.startswith("# "):
            start_idx = i + 1
            break

    for raw in lines[start_idx:]:
        line = raw.rstrip()

        fence_match = re.match(r"^```(\w+)?\s*$", line)
        if fence_match:
            if not in_code_block:
                in_code_block = True
                code_fence_lang = fence_match.group(1) or ""
                code_lines = []
            else:
                # End of block
                flush_code_block()
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if not line.strip():
            doc.add_paragraph("")
            continue

        # Headings
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            continue

        # Bullets
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line[2:].strip())
            continue
        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            p.add_run(re.sub(r"^\d+\.\s+", "", line).strip())
            continue

        # Blockquotes / horizontal rules / markdown tables: keep as plain text
        doc.add_paragraph(line)

    # In case file ends during a code block
    if in_code_block:
        flush_code_block()

    doc.save(docx_path)


if __name__ == "__main__":
    project_root = Path(__file__).resolve().parent
    md_path = project_root / "PROJECT_REPORT.md"
    docx_path = project_root / "PROJECT_REPORT.docx"
    convert(md_path, docx_path)

    print(f"Generated: {docx_path}")

